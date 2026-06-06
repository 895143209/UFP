from __future__ import annotations

import json
import math
import re
from collections import Counter, OrderedDict
from pathlib import Path

import openpyxl

ROOT = Path.cwd()
SECTION_TABLE = ROOT / "\u622a\u9762\u7ba1\u7406\u8868.xlsx"
MODEL_TABLE = ROOT / "\u5efa\u6a21\u7ba1\u7406\u8868.xlsx"
TABLE = MODEL_TABLE if MODEL_TABLE.exists() else SECTION_TABLE
TABLE_DRIVES_MODEL = TABLE == MODEL_TABLE
SECTION_FILE = ROOT / "table_sections.py"
MAIN_NOTEBOOK = ROOT / "Main.ipynb"

HEADER_ROW = 4
DATA_ROW = 5

COL = {
    "kind": 1, "name": 2, "b": 3, "h": 4, "rebar": 5, "stirrup": 6,
    "cover_mat": 7, "core_mat": 8, "rebar_mat": 9, "explicit_rebar_mat": 10,
    "cover": 11, "mesh": 12, "frame_mesh": 13, "column_mesh": 14,
    "wall_zone": 15, "wall_rebar": 16, "edge_h_rebar": 17, "edge_width": 18,
    "edge_rebar": 19, "wall_full_rebar": 20, "wall_x_mesh": 21, "wall_z_mesh": 22,
    "wall_layers": 23, "left_edge_width": 24, "right_edge_width": 25,
    "left_edge_rebar": 26, "right_edge_rebar": 27, "left_edge_h_rebar": 28,
    "right_edge_h_rebar": 29, "full_edge_x": 30, "enabled": 31, "story": 32,
    "order": 33, "func": 34, "length": 35, "anchor_x": 36, "anchor_y": 37,
    "anchor_z": 38, "axis": 39, "plane": 40, "span_dir": 41, "transf": 42,
    "ele_type": 43, "shell_type": 44,
}


def norm(value):
    if value is None:
        return ""
    return str(value).strip()


def is_yes(value):
    return norm(value).lower() in {"yes", "y", "true", "1", "\u662f", "\u542f\u7528"}


def to_int(value, default=None):
    if value is None or value == "":
        return default
    return int(float(value))


def fmt_number(value):
    value = float(value)
    if abs(value - int(value)) < 1e-9:
        return str(int(value))
    return f"{value:g}".replace(".", "p")


def py_num(value):
    if isinstance(value, str):
        s = value.strip()
        try:
            value = float(s)
        except Exception:
            return s
    if value is None:
        return "None"
    value = float(value)
    if abs(value - int(value)) < 1e-9:
        return str(int(value))
    return repr(round(value, 9))


def py_expr(value):
    if value is None or value == "":
        return "None"
    if isinstance(value, str):
        s = value.strip()
        try:
            float(s)
        except Exception:
            return s
    return py_num(value)


def quote(value):
    return repr(str(value))


def sanitize_token(value):
    value = re.sub(r"[^0-9A-Za-z]+", "_", str(value)).strip("_")
    if not value:
        return "X"
    if value[0].isdigit():
        value = "S_" + value
    return value


def material_label(value):
    label = norm(value)
    for prefix in ("TagFrame", "TagWall", "TagSteel", "Tag"):
        if label.startswith(prefix):
            label = label[len(prefix):]
            break
    for suffix in ("Plane", "Plate"):
        if label.endswith(suffix):
            label = label[:-len(suffix)]
    return sanitize_token(label)


def material_expr(value, fallback):
    return norm(value) or fallback


def parse_bar_count_dia(text, default=(0, 0.0)):
    s = norm(text).replace("\u03c6", "\u03a6").replace("D", "\u03a6")
    m = re.search(r"(\d+)\s*[\u03a6@xX*]\s*(\d+(?:\.\d+)?)", s)
    if not m:
        m = re.search(r"(\d+)\D+(\d+(?:\.\d+)?)", s)
    if not m:
        return default
    return int(m.group(1)), float(m.group(2))


def parse_stirrup_dia_spacing(text):
    s = norm(text).replace("\u03c6", "\u03a6").replace("D", "\u03a6")
    m = re.search(r"\u03a6\s*(\d+(?:\.\d+)?)(?:\s*@\s*(\d+(?:\.\d+)?))?", s)
    if not m:
        m = re.search(r"(\d+(?:\.\d+)?)(?:\s*@\s*(\d+(?:\.\d+)?))?", s)
    if not m:
        return 8.0, None
    return float(m.group(1)), None if m.group(2) is None else float(m.group(2))


def parse_frame_bars(row, is_beam=False):
    text = norm(row.get("rebar"))
    if is_beam:
        parts = [p for p in re.split(r"[;\uff1b,\uff0c]", text) if p.strip()]
        top = parse_bar_count_dia(parts[0], (2, 22.0)) if parts else (2, 22.0)
        bottom = parse_bar_count_dia(parts[1], top) if len(parts) > 1 else top
        return top, bottom
    return parse_bar_count_dia(text, (12, 16.0))


def parse_wall_rebar_spec(text):
    s = norm(text)
    vm = re.search(r"V\s*(\d+(?:\.\d+)?)\s*@\s*(\d+(?:\.\d+)?)", s, re.I)
    hm = re.search(r"H\s*(\d+(?:\.\d+)?)\s*@\s*(\d+(?:\.\d+)?)", s, re.I)
    v = (float(vm.group(1)), float(vm.group(2))) if vm else (10.0, 200.0)
    h = (float(hm.group(1)), float(hm.group(2))) if hm else (8.0, 200.0)
    return v, h


def parse_coord_spec(value, length=None):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        n = int(value)
        if length is None:
            return n
        return [round(float(length) * i / n, 6) for i in range(n + 1)]
    s = str(value).strip()
    if re.fullmatch(r"\d+", s):
        n = int(s)
        if length is None:
            return n
        return [round(float(length) * i / n, 6) for i in range(n + 1)]
    parts = [p for p in re.split(r"[;,\uff0c\uff1b\s]+", s) if p]
    return [float(p) for p in parts]


def parse_frame_mesh_arg(value, axis="x"):
    spec = parse_coord_spec(value)
    if spec is None:
        return {}
    if isinstance(spec, int):
        return {"num_ele": spec}
    return {"z_coords" if axis == "z" else "x_coords": spec}


def dict_literal(mapping, indent=0):
    pad = " " * indent
    inner = " " * (indent + 4)
    lines = ["{"]
    for key in sorted(mapping):
        lines.append(f"{inner}{key!r}: {mapping[key]!r},")
    lines.append(pad + "}")
    return "\n".join(lines)


def story_number_from_sheet(sheet_name):
    m = re.search(r"\u7b2c\s*(\d+)\s*\u5c42", sheet_name) or re.search(r"(\d+)", sheet_name)
    return int(m.group(1)) if m else None


def read_rows():
    wb = openpyxl.load_workbook(TABLE, data_only=True)
    rows = []
    for ws in wb.worksheets:
        sn = story_number_from_sheet(ws.title)
        if sn is None:
            continue
        for r in range(DATA_ROW, ws.max_row + 1):
            name = norm(ws.cell(r, COL["name"]).value)
            if not name or not is_yes(ws.cell(r, COL["enabled"]).value):
                continue
            kind_raw = norm(ws.cell(r, COL["kind"]).value)
            if name.startswith("col_") or "\u67f1" in kind_raw:
                kind = "column"
            elif name.startswith("beam_") or "\u6881" in kind_raw:
                kind = "beam"
            elif name.startswith("wall_") or "\u5899" in kind_raw:
                kind = "wall"
            else:
                continue
            row = {"kind_raw": kind_raw, "name": name, "sheet": ws.title}
            for key, col in COL.items():
                if key == "kind":
                    row["excel_kind"] = ws.cell(r, col).value
                else:
                    row[key] = ws.cell(r, col).value
            row["kind"] = kind
            row["story"] = to_int(row.get("story"), sn)
            row["order"] = to_int(row.get("order"), r)
            rows.append(row)
    rows.sort(key=lambda x: (x["story"], x["order"], x["name"]))
    return rows


def frame_section_name(row):
    b, h = row["b"], row["h"]
    cover_mat = material_label(row.get("cover_mat"))
    core_mat = material_label(row.get("core_mat"))
    rebar_mat = material_label(row.get("rebar_mat"))
    stirrup_dia, spacing = parse_stirrup_dia_spacing(row.get("stirrup"))
    if row["kind"] == "column":
        count, dia = parse_frame_bars(row, is_beam=False)
        return sanitize_token(f"Col_{fmt_number(b)}x{fmt_number(h)}_Cvr{cover_mat}_Core{core_mat}_Rb{rebar_mat}_{count}D{fmt_number(dia)}_D{fmt_number(stirrup_dia)}")
    top, bottom = parse_frame_bars(row, is_beam=True)
    sp = f"_{fmt_number(spacing)}" if spacing is not None else ""
    return sanitize_token(f"Beam_{fmt_number(b)}x{fmt_number(h)}_Cvr{cover_mat}_Core{core_mat}_Rb{rebar_mat}_T{top[0]}D{fmt_number(top[1])}_B{bottom[0]}D{fmt_number(bottom[1])}_D{fmt_number(stirrup_dia)}{sp}")


def wall_section_name(row):
    v, hbar = parse_wall_rebar_spec(row.get("wall_rebar"))
    return sanitize_token(f"Wall_t{fmt_number(row['h'])}_Cvr{material_label(row.get('cover_mat'))}_Core{material_label(row.get('core_mat'))}_V{fmt_number(v[0])}_{fmt_number(v[1])}_H{fmt_number(hbar[0])}_{fmt_number(hbar[1])}")


def build_data(rows):
    frame_sections = OrderedDict()
    wall_sections = OrderedDict()
    frame_int_by_comp = {}
    frame_mesh_args = {}
    wall_x = {}
    wall_z = {}
    wall_args = {}
    story_rows = OrderedDict()
    next_frame_sec = 50000
    next_frame_int = 60000
    next_wall_sec = 70000

    for row in rows:
        story_rows.setdefault(row["story"], []).append(row)
        if row["kind"] in {"column", "beam"}:
            key = (row["kind"], float(row["b"]), float(row["h"]), norm(row.get("rebar")), norm(row.get("stirrup")), material_expr(row.get("cover_mat"), "TagFrameC30"), material_expr(row.get("core_mat"), "TagFrameC30Conf45"), material_expr(row.get("rebar_mat"), "TagSteelHRB400"), float(row.get("cover") or 25), float(row.get("mesh") or 50))
            if key not in frame_sections:
                frame_sections[key] = {"row": row, "sec_tag": next_frame_sec, "int_tag": next_frame_int, "name": frame_section_name(row)}
                next_frame_sec += 1
                next_frame_int += 1
            info = frame_sections[key]
            frame_int_by_comp[row["name"]] = info["int_tag"]
            mesh_value = row.get("column_mesh") if row["kind"] == "column" else row.get("frame_mesh")
            frame_mesh_args[row["name"]] = parse_frame_mesh_arg(mesh_value, axis=("z" if row["kind"] == "column" else "x"))
        else:
            v, hbar = parse_wall_rebar_spec(row.get("wall_rebar"))
            key = (float(row["h"]), float(row.get("cover") or 25), int(row.get("wall_layers") or 2), material_expr(row.get("cover_mat"), "TagWallC30Plate"), material_expr(row.get("core_mat"), "TagWallC30Plate"), v, hbar)
            if key not in wall_sections:
                wall_sections[key] = {"row": row, "sec_tag": next_wall_sec, "name": wall_section_name(row), "v": v, "hbar": hbar}
                next_wall_sec += 1
            info = wall_sections[key]
            wall_args[row["name"]] = {"zone_type": "center_only", "sec_tag_center": info["sec_tag"]}
            wall_x[row["name"]] = parse_coord_spec(row.get("wall_x_mesh"), length=float(row["b"]))
            wall_z[row["name"]] = parse_coord_spec(row.get("wall_z_mesh"), length=None)
    return {"rows": rows, "story_rows": story_rows, "frame_sections": frame_sections, "wall_sections": wall_sections, "frame_int_by_comp": frame_int_by_comp, "frame_mesh_args": frame_mesh_args, "wall_x": wall_x, "wall_z": wall_z, "wall_args": wall_args}


def emit_section_file(data):
    lines = []
    lines += [f"# Auto-generated from {TABLE.name} by ufp.apply_model_table().", "# Do not edit by hand; update the workbook and regenerate.", "", "import numpy as np", "", "TABLE_SECTION_SHOW = bool(globals().get('TABLE_SECTION_SHOW', False))", "", "if hasattr(ufp, 'reset_vertical_rebar_registry'):", "    ufp.reset_vertical_rebar_registry()", ""]
    lines += ["def _create_table_wall_section(*, tag, name, thickness, cover, concrete_layers, cover_concrete_tag, core_concrete_tag, h_bar_dia=None, h_bar_spacing=None, v_bar_dia=None, v_bar_spacing=None):", "    h_t = ufp.rebar_equivalent_layer_thickness(h_bar_dia, h_bar_spacing) if h_bar_dia and h_bar_spacing else 0.0", "    v_t = ufp.rebar_equivalent_layer_thickness(v_bar_dia, v_bar_spacing) if v_bar_dia and v_bar_spacing else 0.0", "    core_total = thickness - 2.0 * cover - 2.0 * h_t - 2.0 * v_t", "    if core_total <= 0.0: raise ValueError(f'Layer thickness is not enough for {name}')", "    core_t = core_total / int(concrete_layers)", "    layers = [{'matTag': cover_concrete_tag, 'thickness': cover}]", "    if h_t: layers.append({'matTag': TagreinfH, 'thickness': h_t, 'rebar': True})", "    if v_t: layers.append({'matTag': TagreinfV, 'thickness': v_t, 'rebar': True})", "    for _ in range(int(concrete_layers)): layers.append({'matTag': core_concrete_tag, 'thickness': core_t})", "    if v_t: layers.append({'matTag': TagreinfV, 'thickness': v_t, 'rebar': True})", "    if h_t: layers.append({'matTag': TagreinfH, 'thickness': h_t, 'rebar': True})", "    layers.append({'matTag': cover_concrete_tag, 'thickness': cover})", "    section_data = {'tag': tag, 'name': name, 'thickness': thickness, 'material_properties': {cover_concrete_tag: {'name': str(cover_concrete_tag), 'color': '#E2E764'}, core_concrete_tag: {'name': str(core_concrete_tag), 'color': '#C9D95B'}, TagreinfH: {'name': 'Horizontal PlateRebar', 'color': '#d62728'}, TagreinfV: {'name': 'Vertical PlateRebar', 'color': '#3327d6'}}, 'layers': layers}", "    return ufp.create_layershell(section_data, show=TABLE_SECTION_SHOW)", ""]
    lines += ["def wall_x_coords(component_name, wall_width):", "    coords = WALL_X_COORDS_BY_COMPONENT[component_name]", "    if abs(coords[-1] - float(wall_width)) > 1e-6: raise ValueError(f'{component_name}: wall x coords do not match wall_width={wall_width}')", "    return coords", "", "def wall_z_coords(component_name, wall_height):", "    spec = WALL_Z_SPEC_BY_COMPONENT[component_name]", "    if isinstance(spec, int): return np.linspace(0.0, float(wall_height), spec + 1)", "    coords = list(spec)", "    if abs(coords[-1] - float(wall_height)) > 1e-6: raise ValueError(f'{component_name}: wall z coords do not match wall_height={wall_height}')", "    return coords", ""]
    lines += ["TABLE_SECTIONS_READY = globals().get('TABLE_SECTIONS_READY', False)", "TABLE_SECTION_CREATION_ERROR = None", "", "FRAME_INTEGRATION_BY_COMPONENT = " + dict_literal(data["frame_int_by_comp"]), "FRAME_MESH_ARGS_BY_COMPONENT = " + dict_literal(data["frame_mesh_args"]), "WALL_X_COORDS_BY_COMPONENT = " + dict_literal(data["wall_x"]), "WALL_Z_SPEC_BY_COMPONENT = " + dict_literal(data["wall_z"]), "WALL_SECTION_ARGS_BY_COMPONENT = " + dict_literal(data["wall_args"]), "", "try:", "    if not TABLE_SECTIONS_READY:", "        # -------------------- Frame fiber sections --------------------"]
    for info in data["frame_sections"].values():
        row = info["row"]; var = info["name"]
        lines += [f"        {var}_SecTag = {info['sec_tag']}", f"        {var}_IntTag = {info['int_tag']}", ""]
        if row["kind"] == "column":
            count, dia = parse_frame_bars(row, is_beam=False)
            lines.append(f"        _, {var}_props = ufp.build_rect_section(")
            lines.append(f"            sec_tag={var}_SecTag, b={py_num(row['b'])}, h={py_num(row['h'])}, sec_name={var!r},")
            lines.append(f"            perimeter_bars=({count}, {fmt_number(dia)}), cover={py_num(row.get('cover') or 25)}, stirrup_dia={fmt_number(parse_stirrup_dia_spacing(row.get('stirrup'))[0])},")
        else:
            top, bottom = parse_frame_bars(row, is_beam=True); stirrup_dia, _ = parse_stirrup_dia_spacing(row.get("stirrup"))
            lines.append(f"        _, {var}_props = ufp.build_rect_section(")
            lines.append(f"            sec_tag={var}_SecTag, b={py_num(row['b'])}, h={py_num(row['h'])}, sec_name={var!r},")
            lines.append(f"            top_bars=[({top[0]}, {fmt_number(top[1])})], bottom_bars=[({bottom[0]}, {fmt_number(bottom[1])})], cover={py_num(row.get('cover') or 25)}, stirrup_dia={fmt_number(stirrup_dia)},")
        lines += [f"            mesh_size_cover={py_num(row.get('mesh') or 50)}, mesh_size_core={py_num(row.get('mesh') or 50)},", f"            cover_mat_tag={material_expr(row.get('cover_mat'), 'TagFrameC30')}, core_mat_tag={material_expr(row.get('core_mat'), 'TagFrameC30Conf45')}, rebar_mat_tag={material_expr(row.get('rebar_mat'), 'TagSteelHRB400')},", "            display_results=False, show=TABLE_SECTION_SHOW, register=False,", "        )", f"        ufp.register_frame_section_from_props({var}_SecTag, {var}_props, name={var!r})", f"        ufp.define_frame_integration('Lobatto', {var}_IntTag, {var}_SecTag, 5)", ""]
    lines.append("        # -------------------- Wall layered shell sections --------------------")
    for info in data["wall_sections"].values():
        row = info["row"]; var = info["name"]; v = info["v"]; hbar = info["hbar"]
        lines += [f"        {var}_SecTag = {info['sec_tag']}", f"        _create_table_wall_section(tag={var}_SecTag, name={var!r}, thickness={py_num(row['h'])}, cover={py_num(row.get('cover') or 25)}, concrete_layers={to_int(row.get('wall_layers'), 2)}, cover_concrete_tag={material_expr(row.get('cover_mat'), 'TagWallC30Plate')}, core_concrete_tag={material_expr(row.get('core_mat'), 'TagWallC30Plate')}, h_bar_dia={fmt_number(hbar[0])}, h_bar_spacing={fmt_number(hbar[1])}, v_bar_dia={fmt_number(v[0])}, v_bar_spacing={fmt_number(v[1])})", ""]
    lines += ["        TABLE_SECTIONS_READY = True", "except Exception as exc:", "    TABLE_SECTION_CREATION_ERROR = exc", "    raise", ""]
    SECTION_FILE.write_text("\n".join(lines), encoding="utf-8")


def anchor_point_expr(row):
    return f"({py_expr(row.get('anchor_x'))}, {py_expr(row.get('anchor_y'))}, {py_expr(row.get('anchor_z'))})"


def emit_component(row):
    name = row["name"]
    if row["kind"] == "column":
        return (f"{name} = ufp.column_mesh(\n"
                f"    column_height={py_expr(row.get('length'))},\n"
                f"    anchor_point={anchor_point_expr(row)},\n"
                f"    **FRAME_MESH_ARGS_BY_COMPONENT[{name!r}],\n"
                f"    transf_tag={py_expr(row.get('transf'))},\n"
                f"    integration_tag=FRAME_INTEGRATION_BY_COMPONENT[{name!r}],\n"
                f"    ele_type={quote(row.get('ele_type') or 'forceBeamColumn')},\n"
                f"    node_map=node_map,\n)\nnode_map = {name}[\"node_map\"]\n")
    if row["kind"] == "beam":
        return (f"{name} = ufp.beam_mesh(\n"
                f"    beam_length={py_expr(row.get('length'))},\n"
                f"    anchor_point={anchor_point_expr(row)},\n"
                f"    axis={quote(row.get('axis') or 'X')},\n"
                f"    span_dir={quote(row.get('span_dir') or '+X')},\n"
                f"    **FRAME_MESH_ARGS_BY_COMPONENT[{name!r}],\n"
                f"    transf_tag={py_expr(row.get('transf'))},\n"
                f"    integration_tag=FRAME_INTEGRATION_BY_COMPONENT[{name!r}],\n"
                f"    ele_type={quote(row.get('ele_type') or 'forceBeamColumn')},\n"
                f"    node_map=node_map,\n)\nnode_map = {name}[\"node_map\"]\n")
    return (f"{name} = ufp.wall_mesh_zoned(\n"
            f"    wall_width={py_expr(row.get('b'))},\n"
            f"    wall_height={py_expr(row.get('length'))},\n"
            f"    anchor_point={anchor_point_expr(row)},\n"
            f"    plane={quote(row.get('plane') or 'XOZ')},\n"
            f"    span_dir={quote(row.get('span_dir') or '+X')},\n"
            f"    x_coords=wall_x_coords({name!r}, {py_expr(row.get('b'))}),\n"
            f"    z_coords=wall_z_coords({name!r}, {py_expr(row.get('length'))}),\n"
            f"    **WALL_SECTION_ARGS_BY_COMPONENT[{name!r}],\n"
            f"    shell_type={quote(row.get('shell_type') or 'ShellMITC4')},\n"
            f"    node_map=node_map,\n)\nnode_map = {name}[\"node_map\"]\n")


def generate_story_files(data):
    for story, rows in data["story_rows"].items():
        out = [f"# Auto-generated from {TABLE.name}.", f"# Run from the main notebook with: %run -i model_story_{story:02d}.py", "# This file shares the notebook namespace when executed with %run -i.", "", "# ============================================================", f"# Story {story}", "# ============================================================", "", "# -- table sections guard", "if not globals().get('TABLE_SECTIONS_READY', False):", "    _table_section_error = globals().get('TABLE_SECTION_CREATION_ERROR')", "    _table_section_detail = f' Last table_sections.py error: {_table_section_error}' if _table_section_error else ''", "    raise RuntimeError('table_sections.py did not finish. Run `%run -i table_sections.py` first, or restart the kernel and run Main.ipynb from the top.' + _table_section_detail)", "# -- end table sections guard", ""]
        for title, kind in (("Columns", "column"), ("Walls", "wall"), ("Beams", "beam")):
            group = [r for r in rows if r["kind"] == kind]
            if not group:
                continue
            out += [f"# -------------------- {title} --------------------", ""]
            for row in group:
                out.append(emit_component(row))
        (ROOT / f"model_story_{story:02d}.py").write_text("\n".join(out), encoding="utf-8")


def update_main_notebook():
    return False


def main():
    rows = read_rows()
    data = build_data(rows)
    emit_section_file(data)
    generate_story_files(data)
    changed = update_main_notebook()
    counts = Counter(r["kind"] for r in rows)
    wall_zone_counts = Counter(norm(r.get("wall_zone")) or "center_only" for r in rows if r["kind"] == "wall")
    result = {"table": TABLE.name, "rows": len(rows), "columns": counts.get("column", 0), "beams": counts.get("beam", 0), "walls": counts.get("wall", 0), "frame_sections": len(data["frame_sections"]), "wall_sections": len(data["wall_sections"]), "story_files": len(data["story_rows"]), "main_notebook_changed": changed, "wall_zone_counts": wall_zone_counts}
    printable = dict(result)
    printable["wall_zone_counts"] = dict(wall_zone_counts)
    print(json.dumps(printable, ensure_ascii=False, indent=2))
    return result


if __name__ == "__main__":
    main()
